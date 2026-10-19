from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from src.config import RESUME_DIR, ROOT

NAVY = RGBColor(0x17, 0x3B, 0x57)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
GRAY = RGBColor(0x4B, 0x55, 0x63)
LIGHT = RGBColor(0x6B, 0x72, 0x80)


def _set_font(run, size: float = 8.8, bold: bool = False, color: RGBColor = GRAY) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def _para(container, text: str = "", size: float = 8.8, bold: bool = False, color: RGBColor = GRAY, space_after: float = 1.5, align=None):
    paragraph = container.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.05
    if align is not None:
        paragraph.alignment = align
    if text:
        run = paragraph.add_run(text)
        _set_font(run, size=size, bold=bold, color=color)
    return paragraph


def _section(container, title: str) -> None:
    paragraph = container.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(title)
    _set_font(run, size=10.5, bold=True, color=NAVY)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "0E7C86")
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)


def _bullet(container, text: str, size: float = 8.6, space_after: float = 1.2) -> None:
    paragraph = container.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.02
    paragraph.paragraph_format.left_indent = Inches(0.13)
    paragraph.paragraph_format.first_line_indent = Inches(-0.1)
    run = paragraph.add_run(text)
    _set_font(run, size=size, color=GRAY)


def _bold_prefix(container, prefix: str, text: str, size: float = 8.6) -> None:
    paragraph = container.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1.2)
    paragraph.paragraph_format.line_spacing = 1.02
    paragraph.paragraph_format.left_indent = Inches(0.13)
    paragraph.paragraph_format.first_line_indent = Inches(-0.1)
    run = paragraph.add_run("▸ " + prefix)
    _set_font(run, size=size, bold=True, color=NAVY)
    run2 = paragraph.add_run(text)
    _set_font(run2, size=size, color=GRAY)


def _load_project_metrics() -> dict:
    path = ROOT / "outputs" / "reports" / "project_metrics.json"
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return {
        "data_rows": "3,000 万+",
        "metric_count": "13",
        "eval_questions": "42",
        "sql_accuracy": "85.3%",
        "knowledge_accuracy": "87.5%",
        "guardrail_rate": "100%",
        "report_generation_time": "< 1 分钟",
        "scorecard_auc": "0.726",
        "scorecard_ks": "0.335",
        "insight_count": "5",
    }


def _load_contact() -> dict:
    path = ROOT / "private" / "contact.json"
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return {"email": "[邮箱]", "phone": "[电话]", "location": "[城市]"}


def build_resume() -> Path:
    RESUME_DIR.mkdir(parents=True, exist_ok=True)
    metrics = _load_project_metrics()
    contact = _load_contact()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.36)
    section.left_margin = Inches(0.46)
    section.right_margin = Inches(0.46)

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    left, right = table.rows[0].cells
    left.width = Inches(2.05)
    right.width = Inches(5.18)
    left.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    right.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    # Left column
    photo_path = ROOT / "assets" / "photo.jpg"
    if photo_path.exists():
        photo_para = left.paragraphs[0]
        photo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        photo_para.paragraph_format.space_after = Pt(3)
        photo_para.add_run().add_picture(str(photo_path), width=Inches(0.98))
    else:
        _para(left, "", size=2)

    _section(left, "联系方式")
    for item in [
        f"邮箱：{contact['email']}",
        f"电话：{contact['phone']}",
        f"籍贯：{contact['location']}",
    ]:
        _para(left, item, size=8.3, space_after=1)

    _section(left, "专业技能")
    skill_groups = [
        ("数据工具", "SQL（MySQL/DuckDB）、Python（Pandas/NumPy/scikit-learn）、Excel（Power Query/透视表）、Power BI/FineBI"),
        ("AI 工程", "Prompt 工程、DeepSeek/OpenAI 兼容 API、NL2SQL、RAG、Agent 工作流、Guardrails、评测集设计"),
        ("分析能力", "漏斗/cohort/vintage、客户分层、归因分解、异常检测、WOE/IV 评分卡"),
        ("产品落地", "PRD、用户故事、UAT、培训、采纳度量、跨部门协同"),
    ]
    for title, body in skill_groups:
        _para(left, title, size=8.3, bold=True, color=TEAL, space_after=0.5)
        _para(left, body, size=7.9, space_after=2.2)

    _section(left, "教育背景")
    _para(left, "中国民用航空飞行学院", size=8.4, bold=True, color=NAVY, space_after=0.5)
    _para(left, "交通信息工程及控制 · 学术硕士", size=7.9, space_after=0.3)
    _para(left, "2021.09 – 2024.06", size=7.7, color=LIGHT, space_after=2)
    _para(left, "四川农业大学（211）", size=8.4, bold=True, color=NAVY, space_after=0.5)
    _para(left, "生物工程 · 本科", size=7.9, space_after=0.3)
    _para(left, "2016.09 – 2020.06", size=7.7, color=LIGHT, space_after=2)

    _section(left, "荣誉奖项")
    for item in ["核心期刊论文 3 篇", "研究生电子设计竞赛三等奖", "网络安全攻防实战优秀个人", "研究生二等奖学金"]:
        _bullet(left, item, size=7.8, space_after=0.6)

    # Right column
    name = right.paragraphs[0]
    name.paragraph_format.space_after = Pt(0)
    run = name.add_run("魏嗣杰")
    _set_font(run, size=21, bold=True, color=NAVY)
    run2 = name.add_run("  数据分析师（AI 方向）")
    _set_font(run2, size=10.5, bold=True, color=TEAL)

    _section(right, "个人简介")
    _para(
        right,
        "1 年+消费金融数据分析经验，熟练使用 SQL / Python / Excel / Power BI 完成全链路分析。"
        "擅长漏斗转化、客户分层、监控看板与指标口径治理，能将业务问题拆解为可落地的数据方案。"
        "近期以 Home Credit 公开数据独立完成消费金融经营分析 AI Copilot，覆盖申请—审批—用信—还款—逾期全链路，"
        "并建立 NL2SQL 语义层、RAG 知识库与评测体系。",
        size=8.5,
        space_after=2,
    )

    _section(right, "工作经历")
    job = right.add_paragraph()
    job.paragraph_format.space_after = Pt(0.5)
    r1 = job.add_run("兴业消费金融有限公司")
    _set_font(r1, size=9.6, bold=True, color=NAVY)
    r2 = job.add_run("  数据分析师")
    _set_font(r2, size=8.8, bold=True, color=TEAL)
    r3 = job.add_run("  2025.03 – 2026.04")
    _set_font(r3, size=8.0, color=LIGHT)
    _bold_prefix(right, "全渠道获客转化分析：", "主导获客—进件—审批—放款全链路分析，拆解转化漏斗定位 3 个关键流失节点；推动中介放款占比下降 12%、支行直客占比提升 14%，累计带动放款超 1200 万元，咨询转化率 48%。")
    _bold_prefix(right, "客户运营与风险监控：", "清洗存量客户行为与还款数据，搭建 3 大类 12 项信贷运营监控看板，实现转化下滑与客群逾期异常提前预警；保障放款成功率 99.5%、不良率 <0.3%；高价值客户复贷率提升至 35%。")
    _bold_prefix(right, "数据效能与团队赋能：", "搭建多维度自动校验体系与标准化分析模板，报告撰写周期缩短 40%，数据准确率 100%、报告通过率 92%；沉淀团队统一数据报表体系与工具包。")
    _bold_prefix(right, "工作呼应：", "将指标口径字典、异常检测与自动化报告思路复用于经营报表和专项分析，降低口径争议与人工核对成本。")

    _section(right, "项目经历")
    project = right.add_paragraph()
    project.paragraph_format.space_after = Pt(0.5)
    p1 = project.add_run("消费金融经营分析 AI Copilot")
    _set_font(p1, size=9.6, bold=True, color=NAVY)
    p2 = project.add_run("  个人项目")
    _set_font(p2, size=8.2, bold=True, color=TEAL)
    p3 = project.add_run("  2026.09 – 2026.10")
    _set_font(p3, size=8.0, color=LIGHT)
    _bold_prefix(right, "数据底座：", f"基于 Home Credit 7 张表（约 {metrics['data_rows']} 行）构建 DuckDB 星型模型，定义 {metrics['metric_count']} 个指标口径，覆盖审批漏斗、组合结构、vintage 迁徙与客户分层。")
    _bold_prefix(right, "经营诊断：", f"完成渠道/产品/地区/客群四维分析，使用结构效应与比率效应分解定位审批通过率变化来源；构建 WOE/IV 风险评分卡（AUC {metrics['scorecard_auc']}、KS {metrics['scorecard_ks']}），PD 校准后风险分层单调。")
    _bold_prefix(right, "AI Copilot：", f"设计语义层与只读 NL2SQL 链路，DeepSeek 生成 SQL 后经 sqlglot 校验、DuckDB 只读执行与结果复核；构建 RAG 知识库与 {metrics['eval_questions']} 题评测集，SQL 准确率 {metrics['sql_accuracy']}、知识问答 {metrics['knowledge_accuracy']}，危险 SQL 拦截率 {metrics['guardrail_rate']}。")
    _bold_prefix(right, "产品化：", f"产出 PRD、UAT、培训手册与推广方案；HTML 看板与 Excel 工作簿支持经营例会，单次月报生成 {metrics['report_generation_time']}（本机实测）。")

    output = RESUME_DIR / "魏嗣杰_数据分析师_AI方向_一页版.docx"
    doc.save(output)
    return output


if __name__ == "__main__":
    print(build_resume())
