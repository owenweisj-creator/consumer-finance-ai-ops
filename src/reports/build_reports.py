from __future__ import annotations

import json
import re
import subprocess
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import duckdb
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from src.ai.llm import DeepSeekClient, LLMError
from src.analytics.metrics import anomaly_scan, connect, query_df, rate_decomposition
from src.config import DB_PATH, PROMPTS_DIR, REPORTS_DIR
from src.reports.charts import channel_chart, funnel_chart, kpi_trend_chart, vintage_heatmap

NAVY = RGBColor(0x17, 0x3B, 0x57)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
GRAY = RGBColor(0x6B, 0x72, 0x80)


def _fmt(value: Any, digits: int = 0) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return "-"
    if digits == 0:
        return f"{float(value):,.0f}"
    return f"{float(value):,.{digits}f}"


def _pct(value: Any, digits: int = 2) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return "-"
    return f"{float(value) * 100:.{digits}f}%"


def _signed_pct(value: Any, digits: int = 1) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return "-"
    return f"{float(value) * 100:+.{digits}f}%"


def monthly_facts(period: str) -> dict[str, Any]:
    con = connect()
    try:
        current = con.execute(
            "SELECT * FROM mart.v_kpi_monthly WHERE year_month = ?", [period]
        ).df()
        if current.empty:
            raise ValueError(f"No KPI data for period {period}")
        latest = current.iloc[0].to_dict()
        previous_period = con.execute(
            "SELECT max(year_month) FROM mart.v_kpi_monthly WHERE year_month < ?", [period]
        ).fetchone()[0]
        previous = con.execute(
            "SELECT * FROM mart.v_kpi_monthly WHERE year_month = ?", [previous_period]
        ).df().iloc[0].to_dict()
        channels = con.execute(
            """
            SELECT channel_type, sum(applications) AS applications, sum(approved) AS approved,
                   round(sum(approved)/NULLIF(sum(applications),0),4) AS approval_rate,
                   round(sum(used)/NULLIF(sum(approved),0),4) AS utilization_rate,
                   sum(approved_amount) AS approved_amount
            FROM mart.v_funnel_monthly WHERE year_month = ? GROUP BY 1 ORDER BY applications DESC
            """,
            [period],
        ).df()
        products = con.execute(
            """
            SELECT contract_type, product_type, sum(applications) AS applications,
                   round(sum(approved)/NULLIF(sum(applications),0),4) AS approval_rate,
                   sum(approved_amount) AS approved_amount
            FROM mart.v_funnel_monthly WHERE year_month = ? GROUP BY 1,2 ORDER BY applications DESC
            """,
            [period],
        ).df()
        regions = con.execute("SELECT * FROM mart.v_region_monthly ORDER BY region_rating").df()
        vintage = con.execute("SELECT * FROM mart.v_vintage ORDER BY cohort_month, mob").df()
        trend = con.execute(
            """
            SELECT * FROM mart.v_kpi_monthly
            WHERE month_start >= (SELECT max(month_start) - INTERVAL '11 months' FROM mart.v_kpi_monthly)
            ORDER BY month_start
            """
        ).df()
    finally:
        con.close()

    decomposition = rate_decomposition("channel_type", periods=2)
    alerts = [item for item in anomaly_scan() if item["period"] == period]
    return {
        "report_type": "monthly",
        "period": period,
        "anchor_date": "2026-09-30",
        "current": latest,
        "previous": previous,
        "previous_period": previous_period,
        "channels": channels.to_dict(orient="records"),
        "products": products.to_dict(orient="records"),
        "regions": regions.to_dict(orient="records"),
        "vintage": vintage.tail(80).to_dict(orient="records"),
        "trend": trend.to_dict(orient="records"),
        "decomposition": decomposition.to_dict(orient="records") if not decomposition.empty else [],
        "alerts": alerts,
    }


def weekly_facts(start: date, end: date) -> dict[str, Any]:
    con = connect()
    try:
        def summary(a: date, b: date) -> dict:
            row = con.execute(
                """
                SELECT count(*) AS applications,
                       sum(CASE WHEN status IN ('Approved','Unused offer') THEN 1 ELSE 0 END) AS approved,
                       sum(CASE WHEN status='Approved' THEN 1 ELSE 0 END) AS used,
                       round(sum(CASE WHEN status IN ('Approved','Unused offer') THEN 1 ELSE 0 END)/NULLIF(count(*),0),4) AS approval_rate,
                       round(sum(CASE WHEN status='Approved' THEN 1 ELSE 0 END)/NULLIF(sum(CASE WHEN status IN ('Approved','Unused offer') THEN 1 ELSE 0 END),0),4) AS utilization_rate,
                       sum(CASE WHEN status='Approved' THEN approved_amount ELSE 0 END) AS approved_amount,
                       round(sum(CASE WHEN status='Approved' THEN approved_amount ELSE 0 END)/NULLIF(sum(CASE WHEN status='Approved' THEN 1 ELSE 0 END),0),2) AS avg_ticket
                FROM mart.fact_previous_application
                WHERE decision_date BETWEEN ? AND ?
                """,
                [a, b],
            ).df().iloc[0].to_dict()
            return row

        current = summary(start, end)
        previous_start = start - timedelta(days=7)
        previous_end = start - timedelta(days=1)
        previous = summary(previous_start, previous_end)
        channels = con.execute(
            """
            SELECT channel_type, count(*) AS applications,
                   round(sum(CASE WHEN status IN ('Approved','Unused offer') THEN 1 ELSE 0 END)/NULLIF(count(*),0),4) AS approval_rate,
                   sum(CASE WHEN status='Approved' THEN approved_amount ELSE 0 END) AS approved_amount
            FROM mart.fact_previous_application
            WHERE decision_date BETWEEN ? AND ?
            GROUP BY 1 ORDER BY applications DESC
            """,
            [start, end],
        ).df()
    finally:
        con.close()
    return {
        "report_type": "weekly",
        "period": f"{start.isoformat()} ~ {end.isoformat()}",
        "current": current,
        "previous": previous,
        "channels": channels.to_dict(orient="records"),
        "alerts": [],
    }


def special_facts(period: str) -> dict[str, Any]:
    decomp = rate_decomposition("channel_type", periods=2)
    if decomp.empty:
        raise ValueError("No decomposition data")
    current = decomp.iloc[0]["current_month"]
    overview = query_df(
        "SELECT * FROM mart.v_kpi_monthly WHERE year_month IN (?, ?) ORDER BY year_month",
        [decomp.iloc[0]["previous_month"], current],
    )
    return {
        "report_type": "special",
        "period": current,
        "topic": "审批通过率变化的结构效应与比率效应分解",
        "overview": overview.to_dict(orient="records"),
        "decomposition": decomp.to_dict(orient="records"),
        "mix_total": round(float(decomp["mix_effect"].sum()), 6),
        "rate_total": round(float(decomp["rate_effect"].sum()), 6),
        "total_effect": round(float(decomp["total_effect"].sum()), 6),
    }


def _fact_numbers(facts: dict[str, Any]) -> list[float]:
    numbers: list[float] = []
    for value in facts.values():
        containers = [value] if isinstance(value, dict) else value if isinstance(value, list) else []
        for container in containers:
            if isinstance(container, dict):
                numbers.extend(
                    float(nested)
                    for nested in container.values()
                    if isinstance(nested, (int, float)) and not isinstance(nested, bool)
                )
    return numbers


def _numeric_candidates(numbers: list[float]) -> set[float]:
    candidates: set[float] = set(numbers)
    for left in numbers:
        candidates.add(left * 100)
        for right in numbers:
            candidates.add(left - right)
            if right:
                candidates.add(left / right)
                candidates.add((left - right) / right * 100)
    return candidates


def _narrative_is_grounded(text: str, facts: dict[str, Any]) -> bool:
    numbers = _fact_numbers(facts)
    candidates = _numeric_candidates(numbers)
    total = 0
    suspicious = 0
    for token in re.findall(r"\d[\d,]*\.?\d*%?", text):
        total += 1
        clean = token.rstrip("%").replace(",", "")
        try:
            numeric = float(clean)
        except ValueError:
            suspicious += 1
            continue
        if numeric in {2024.0, 2025.0, 2026.0}:
            continue
        if abs(numeric) <= 20:
            continue
        tolerance = max(0.02, abs(numeric) * 0.01)
        if any(abs(numeric - candidate) <= tolerance for candidate in candidates):
            continue
        suspicious += 1
    if total == 0:
        return True
    return suspicious <= 5 and (1 - suspicious / total) >= 0.85


def _localize_terms(text: str) -> str:
    replacements = {
        "mix_effect": "结构效应",
        "rate_effect": "比率效应",
        "approval_rate": "审批通过率",
        "utilization_rate": "用信率",
        "unused_offer": "未使用额度",
        "approved_amount": "审批金额",
        "avg_ticket": "件均",
        "pct": "个百分点",
        "p.p.": "个百分点",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def fallback_narrative(facts: dict[str, Any]) -> str:
    if facts["report_type"] == "monthly":
        cur = facts["current"]
        prev = facts["previous"]
        return (
            f"本期进件量 {_fmt(cur.get('applications'))} 笔，较上期 {_signed_pct((cur.get('applications', 0) - prev.get('applications', 0)) / prev.get('applications', 1) if prev.get('applications') else None)}；"
            f"审批通过率 {_pct(cur.get('approval_rate'))}，用信率 {_pct(cur.get('utilization_rate'))}。"
            f"审批金额 {_fmt(cur.get('approved_amount'))} 万元，件均 {_fmt(cur.get('avg_ticket'), 2)} 万元。"
            "建议优先拆解渠道结构变化，再判断同一渠道内的审批策略和客群质量是否发生变化。"
        )
    if facts["report_type"] == "weekly":
        cur = facts["current"]
        return (
            f"本周进件量 {_fmt(cur.get('applications'))} 笔，审批通过率 {_pct(cur.get('approval_rate'))}，"
            f"用信率 {_pct(cur.get('utilization_rate'))}，审批金额 {_fmt(cur.get('approved_amount'))} 万元。"
            "建议关注周环比波动最大的渠道，并核对是否为节假日、活动或审批策略调整导致。"
        )
    return (
        "本期审批通过率变化需要同时看结构效应和比率效应。"
        "若结构效应占主导，应优先调整渠道组合；若比率效应占主导，应进一步核查审批规则和客群质量。"
    )


def generate_narrative(facts: dict[str, Any], client: DeepSeekClient | None = None) -> str:
    audit: dict[str, Any] = {"mode": "fallback", "reason": "unknown"}
    if client is None:
        try:
            client = DeepSeekClient()
        except LLMError:
            audit["reason"] = "missing_api_key"
            (REPORTS_DIR / "narrative_audit.json").write_text(
                json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            return fallback_narrative(facts)
    prompt = (PROMPTS_DIR / "report_v1.md").read_text(encoding="utf-8").replace(
        "{{FACTS}}", json.dumps(facts, ensure_ascii=False, default=str, indent=2)
    )
    try:
        response = client.chat(
            [
                {"role": "system", "content": prompt},
                {"role": "user", "content": "请生成经营解读。"},
            ],
            temperature=0.2,
            max_tokens=1400,
        )
        if _narrative_is_grounded(response.content, facts):
            localized = _localize_terms(response.content)
            audit.update({"mode": "llm", "reason": "grounded", "preview": localized[:300]})
            (REPORTS_DIR / "narrative_audit.json").write_text(
                json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            return localized.strip()
        audit["reason"] = "numeric_grounding_failed"
        audit["preview"] = response.content[:300]
    except Exception as exc:  # noqa: BLE001
        audit["reason"] = f"llm_error: {exc}"[:400]
    (REPORTS_DIR / "narrative_audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return fallback_narrative(facts)


def _set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def _add_table(document: Document, df: pd.DataFrame, columns: dict[str, str], widths: list[float] | None = None) -> None:
    if df.empty:
        document.add_paragraph("暂无数据。")
        return
    table = document.add_table(rows=1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    for index, label in enumerate(columns.values()):
        header[index].text = label
        for paragraph in header[index].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _set_cell_shading(header[index], "173B57")
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for index, key in enumerate(columns):
            value = row.get(key, "")
            cells[index].text = str(value)
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)


def _add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = NAVY if level == 1 else TEAL
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    run.font.size = Pt(10)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _add_rich_paragraph(document: Document, text: str, bullet: bool = False, bold_prefix: bool = False) -> None:
    paragraph = document.add_paragraph(style="List Bullet" if bullet else None)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.22
    if bullet:
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    parts = text.split("**")
    for index, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        is_bold = bold_prefix or index % 2 == 1
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(9.6 if bold_prefix else 10)
        run.bold = is_bold
        run.font.color.rgb = NAVY if is_bold else RGBColor(0x33, 0x3D, 0x4A)


def _add_markdown(document: Document, text: str) -> None:
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("## "):
            _add_rich_paragraph(document, line[3:].strip(), bold_prefix=True)
        elif line.startswith("### "):
            _add_rich_paragraph(document, line[4:].strip(), bold_prefix=True)
        elif line.startswith("- "):
            _add_rich_paragraph(document, line[2:].strip(), bullet=True)
        elif re.match(r"^\d+[.、]\s*", line):
            _add_rich_paragraph(document, line, bullet=True)
        else:
            _add_rich_paragraph(document, line)


def build_monthly_report(period: str) -> dict[str, Path]:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    facts = monthly_facts(period)
    narrative = generate_narrative(facts)
    cur = facts["current"]
    prev = facts["previous"]

    trend_df = pd.DataFrame(facts["trend"])
    channels = pd.DataFrame(facts["channels"])
    funnel_df = pd.DataFrame([cur])
    vintage_df = pd.DataFrame(facts["vintage"])
    kpi_chart = kpi_trend_chart(trend_df, REPORTS_DIR / "chart_kpi_trend.png")
    funnel_img = funnel_chart(funnel_df, REPORTS_DIR / "chart_funnel.png")
    vintage_img = vintage_heatmap(vintage_df, REPORTS_DIR / "chart_vintage.png")
    channel_img = channel_chart(channels, REPORTS_DIR / "chart_channel.png")

    md = [
        f"# 消费金融经营分析月报（{period}）",
        "",
        f"> 数据锚点：2026-09-30｜生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}｜数据源：Home Credit 公开数据",
        "",
        "## 一、经营结论",
        "",
        narrative,
        "",
        "## 二、核心指标",
        "",
        "| 指标 | 本期 | 上期 | 环比 |",
        "|---|---:|---:|---:|",
        f"| 进件量 | {_fmt(cur.get('applications'))} | {_fmt(prev.get('applications'))} | {_signed_pct((cur['applications']-prev['applications'])/prev['applications'] if prev.get('applications') else None)} |",
        f"| 审批通过率 | {_pct(cur.get('approval_rate'))} | {_pct(prev.get('approval_rate'))} | {_signed_pct((cur['approval_rate']-prev['approval_rate'])/prev['approval_rate'] if prev.get('approval_rate') else None)} |",
        f"| 用信率 | {_pct(cur.get('utilization_rate'))} | {_pct(prev.get('utilization_rate'))} | {_signed_pct((cur['utilization_rate']-prev['utilization_rate'])/prev['utilization_rate'] if prev.get('utilization_rate') else None)} |",
        f"| 审批金额 | {_fmt(cur.get('approved_amount'))} | {_fmt(prev.get('approved_amount'))} | {_signed_pct((cur['approved_amount']-prev['approved_amount'])/prev['approved_amount'] if prev.get('approved_amount') else None)} |",
        f"| 件均 | {_fmt(cur.get('avg_ticket'), 2)} | {_fmt(prev.get('avg_ticket'), 2)} | {_signed_pct((cur['avg_ticket']-prev['avg_ticket'])/prev['avg_ticket'] if prev.get('avg_ticket') else None)} |",
        "",
        "## 三、渠道与产品",
        "",
        "### 渠道表现",
        "",
        "| 渠道 | 进件量 | 审批通过率 | 用信率 | 审批金额 |",
        "|---|---:|---:|---:|---:|",
    ]
    for row in facts["channels"]:
        md.append(f"| {row['channel_type']} | {_fmt(row['applications'])} | {_pct(row['approval_rate'])} | {_pct(row['utilization_rate'])} | {_fmt(row['approved_amount'])} |")
    md += ["", "### 产品结构", "", "| 合同类型 | 产品类型 | 进件量 | 审批通过率 | 审批金额 |", "|---|---|---:|---:|---:|"]
    for row in facts["products"]:
        md.append(f"| {row['contract_type']} | {row['product_type']} | {_fmt(row['applications'])} | {_pct(row['approval_rate'])} | {_fmt(row['approved_amount'])} |")

    md += ["", "## 四、风险与 Vintage", "", "Vintage 用于控制账龄结构变化，避免把新客占比上升误判为风险恶化。"]
    if facts["vintage"]:
        last = facts["vintage"][-1]
        md.append(f"最近 cohort 的累计 DPD30 为 {_pct(last.get('cum_dpd30_rate'))}，累计 DPD60 为 {_pct(last.get('cum_dpd60_rate'))}。")
    md += ["", "## 五、异常与归因", ""]
    if facts["alerts"]:
        for alert in facts["alerts"]:
            md.append(f"- {alert['label']}在 {alert['period']} 出现异常（robust z = {alert['robust_z']}），本期 {alert['value']:.4f}。")
    else:
        md.append("- 本月未触发统计异常阈值，但仍需结合业务事件复核。")
    md += ["", "## 六、行动建议", "", "1. 对通过率下降且进件量上升的渠道做专项复核，区分审批规则变化与客群质量变化。", "2. 对用信率偏低的渠道检查额度、定价和触达节奏，避免“批而不用”。", "3. 对高风险地区执行名单制管理，先控额、再提价，最后考虑退出。", "", "## 七、数据说明", "", "- 数据为 Home Credit 公开消费金融数据，业务日历为映射后的模拟时间。", "- 金额单位为万元（归一化），仅用于结构和趋势分析。", "- TARGET 表示 30+ 天逾期标签，不等同于核销率。", "- 本报告由 AI 辅助生成，所有数字均来自项目数据仓库，发布前需人工复核。"]

    md_path = REPORTS_DIR / f"monthly_{period}.md"
    md_path.write_text("\n".join(md), encoding="utf-8")

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    title = document.add_heading(f"消费金融经营分析月报（{period}）", level=0)
    for run in title.runs:
        run.font.color.rgb = NAVY
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    _add_body(document, "数据锚点：2026-09-30｜数据源：Home Credit 公开数据｜AI 辅助生成，发布前需人工复核")

    _add_heading(document, "一、经营结论")
    _add_markdown(document, narrative)
    _add_heading(document, "二、核心指标")
    kpi_table = pd.DataFrame(
        [
            {
                "指标": "进件量",
                "本期": _fmt(cur.get("applications")),
                "上期": _fmt(prev.get("applications")),
                "环比": _signed_pct((cur["applications"] - prev["applications"]) / prev["applications"] if prev.get("applications") else None),
            },
            {
                "指标": "审批通过率",
                "本期": _pct(cur.get("approval_rate")),
                "上期": _pct(prev.get("approval_rate")),
                "环比": _signed_pct((cur["approval_rate"] - prev["approval_rate"]) / prev["approval_rate"] if prev.get("approval_rate") else None),
            },
            {
                "指标": "用信率",
                "本期": _pct(cur.get("utilization_rate")),
                "上期": _pct(prev.get("utilization_rate")),
                "环比": _signed_pct((cur["utilization_rate"] - prev["utilization_rate"]) / prev["utilization_rate"] if prev.get("utilization_rate") else None),
            },
            {
                "指标": "审批金额",
                "本期": _fmt(cur.get("approved_amount")),
                "上期": _fmt(prev.get("approved_amount")),
                "环比": _signed_pct((cur["approved_amount"] - prev["approved_amount"]) / prev["approved_amount"] if prev.get("approved_amount") else None),
            },
        ]
    )
    _add_table(document, kpi_table, {c: c for c in kpi_table.columns})
    document.add_picture(str(kpi_chart), width=Inches(6.4))
    _add_heading(document, "三、渠道与产品")
    channel_table = channels.head(8).copy()
    channel_table["applications"] = channel_table["applications"].map(lambda x: _fmt(x))
    channel_table["approved_amount"] = channel_table["approved_amount"].map(lambda x: _fmt(x, 2))
    channel_table["approval_rate"] = channel_table["approval_rate"].map(lambda x: _pct(x))
    channel_table["utilization_rate"] = channel_table["utilization_rate"].map(lambda x: _pct(x))
    _add_table(
        document,
        channel_table,
        {
            "channel_type": "渠道",
            "applications": "进件量",
            "approval_rate": "通过率",
            "utilization_rate": "用信率",
            "approved_amount": "审批金额",
        },
    )
    document.add_picture(str(channel_img), width=Inches(6.0))
    _add_heading(document, "四、风险与 Vintage")
    if vintage_img.exists():
        document.add_picture(str(vintage_img), width=Inches(6.6))
    _add_heading(document, "五、行动建议")
    for item in [
        "对通过率下降且进件量上升的渠道做专项复核，区分审批规则变化与客群质量变化。",
        "对用信率偏低的渠道检查额度、定价和触达节奏，避免“批而不用”。",
        "对高风险地区执行名单制管理，先控额、再提价，最后考虑退出。",
    ]:
        _add_body(document, "• " + item)
    _add_heading(document, "六、数据说明")
    for item in [
        "数据为 Home Credit 公开消费金融数据，业务日历为映射后的模拟时间。",
        "金额单位为万元（归一化），仅用于结构和趋势分析。",
        "TARGET 表示 30+ 天逾期标签，不等同于核销率。",
        "本报告由 AI 辅助生成，所有数字均来自项目数据仓库，发布前需人工复核。",
    ]:
        _add_body(document, "• " + item)

    docx_path = REPORTS_DIR / f"monthly_{period}.docx"
    document.save(docx_path)
    pdf_path = convert_docx_to_pdf(docx_path)
    return {"markdown": md_path, "docx": docx_path, "pdf": pdf_path, "chart": kpi_chart}


def convert_docx_to_pdf(docx_path: Path) -> Path:
    pdf_path = docx_path.with_suffix(".pdf")
    script = (
        "$ErrorActionPreference='Stop'; "
        "$word=New-Object -ComObject Word.Application; $word.Visible=$false; "
        f"$doc=$word.Documents.Open('{docx_path}'); $doc.SaveAs([ref]'{pdf_path}',[ref]17); "
        "$doc.Close(); $word.Quit()"
    )
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True, capture_output=True)
        return pdf_path
    except Exception:
        return Path("")


def build_weekly_report(start: date, end: date) -> dict[str, Path]:
    facts = weekly_facts(start, end)
    narrative = generate_narrative(facts)
    cur, prev = facts["current"], facts["previous"]
    md = [
        f"# 消费金融经营分析周报（{start.isoformat()} 至 {end.isoformat()}）",
        "",
        f"> 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}｜数据源：Home Credit 公开数据",
        "",
        "## 一、本周结论",
        "",
        narrative,
        "",
        "## 二、核心指标",
        "",
        "| 指标 | 本周 | 上周 | 环比 |",
        "|---|---:|---:|---:|",
        f"| 进件量 | {_fmt(cur.get('applications'))} | {_fmt(prev.get('applications'))} | {_signed_pct((cur['applications']-prev['applications'])/prev['applications'] if prev.get('applications') else None)} |",
        f"| 审批通过率 | {_pct(cur.get('approval_rate'))} | {_pct(prev.get('approval_rate'))} | {_signed_pct((cur['approval_rate']-prev['approval_rate'])/prev['approval_rate'] if prev.get('approval_rate') else None)} |",
        f"| 用信率 | {_pct(cur.get('utilization_rate'))} | {_pct(prev.get('utilization_rate'))} | {_signed_pct((cur['utilization_rate']-prev['utilization_rate'])/prev['utilization_rate'] if prev.get('utilization_rate') else None)} |",
        f"| 审批金额 | {_fmt(cur.get('approved_amount'))} | {_fmt(prev.get('approved_amount'))} | {_signed_pct((cur['approved_amount']-prev['approved_amount'])/prev['approved_amount'] if prev.get('approved_amount') else None)} |",
        "",
        "## 三、渠道表现",
        "",
        "| 渠道 | 进件量 | 审批通过率 | 审批金额 |",
        "|---|---:|---:|---:|",
    ]
    for row in facts["channels"]:
        md.append(f"| {row['channel_type']} | {_fmt(row['applications'])} | {_pct(row['approval_rate'])} | {_fmt(row['approved_amount'])} |")
    md += ["", "## 四、下周行动项", "", "1. 复核周环比波动最大的渠道。", "2. 核对审批规则、活动和节假日的业务事件。", "3. 对异常指标建立行动项并指定负责人。"]
    path = REPORTS_DIR / f"weekly_{start.isoformat()}_{end.isoformat()}.md"
    path.write_text("\n".join(md), encoding="utf-8")
    return {"markdown": path}


def build_special_report(period: str) -> dict[str, Path]:
    facts = special_facts(period)
    decomp = pd.DataFrame(facts["decomposition"])
    total_effect = facts["total_effect"]
    mix_total = facts["mix_total"]
    rate_total = facts["rate_total"]
    top_negative = decomp[decomp["total_effect"] < 0].sort_values("total_effect").head(3)
    top_positive = decomp[decomp["total_effect"] > 0].sort_values("total_effect", ascending=False).head(3)
    narrative = generate_narrative(facts)
    md = [
        f"# 专项分析：审批通过率变化的结构分解（{period}）",
        "",
        "## 结论",
        "",
        narrative,
        "",
        "## 分解方法",
        "",
        "整体通过率变化 = 结构效应 + 比率效应。结构效应用渠道占比变化乘以渠道基期通过率与整体通过率之差；比率效应用当期渠道占比乘以渠道通过率变化。",
        "",
        "## 渠道贡献",
        "",
        "| 渠道 | 基期占比 | 当期占比 | 结构效应 | 比率效应 | 合计 |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for _, row in decomp.sort_values("total_effect", ascending=False).iterrows():
        md.append(
            f"| {row['channel_type']} | {_pct(row['weight_previous'])} | {_pct(row['weight_current'])} | "
            f"{row['mix_effect']:+.4f} | {row['rate_effect']:+.4f} | {row['total_effect']:+.4f} |"
        )
    md += [
        "",
        f"合计效应：结构效应 {mix_total:+.4f}，比率效应 {rate_total:+.4f}，总效应 {total_effect:+.4f}。",
    ]
    md += ["", "## 重点渠道", ""]
    for _, row in top_negative.iterrows():
        md.append(f"- 负向贡献：{row['channel_type']}，合计 {row['total_effect']:+.4f}，其中结构 {row['mix_effect']:+.4f}、比率 {row['rate_effect']:+.4f}。")
    if top_positive.empty:
        md.append("- 正向贡献：本月没有合计效应为正的渠道，说明结构收缩是普遍现象。")
    else:
        for _, row in top_positive.iterrows():
            md.append(f"- 正向贡献：{row['channel_type']}，合计 {row['total_effect']:+.4f}，其中结构 {row['mix_effect']:+.4f}、比率 {row['rate_effect']:+.4f}。")
    md += ["", "## 建议", "", "1. 若结构效应为主，优先优化渠道组合和资源投放。", "2. 若比率效应为主，优先核查审批规则、客群质量和反欺诈策略。", "3. 对负向贡献最大的渠道做一周内可验证的整改实验，并设置通过率与不良率双门槛。"]
    path = REPORTS_DIR / f"special_approval_decomposition_{period}.md"
    path.write_text("\n".join(md), encoding="utf-8")
    return {"markdown": path}


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("--type", choices=["monthly", "weekly", "special"], default="monthly")
    parser.add_argument("--period", default="2026-09")
    parser.add_argument("--start", default="2026-09-21")
    parser.add_argument("--end", default="2026-09-27")
    args = parser.parse_args()
    if args.type == "monthly":
        print(build_monthly_report(args.period))
    elif args.type == "weekly":
        print(build_weekly_report(date.fromisoformat(args.start), date.fromisoformat(args.end)))
    else:
        print(build_special_report(args.period))
